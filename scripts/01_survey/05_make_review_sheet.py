#!/usr/bin/env python3
"""Build the sheet for manual researcher review: packet summary + my labels + LLM verdicts (v2) + empty review columns.
Usage: python3 05_make_review_sheet.py --pred results/design_llm/labeled_v2.opus.pred.jsonl --out results/design_llm/review_sheet_v2.xlsx [--sample 60 --seed 1]"""
import argparse, csv, json, random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
ap = argparse.ArgumentParser(); ap.add_argument("--pred", required=True); ap.add_argument("--out", required=True); ap.add_argument("--sample", type=int, default=0); ap.add_argument("--seed", type=int, default=1); a = ap.parse_args()
packets = {json.loads(l)["study"]: json.loads(l) for l in open("data/design_packets/labeled_v1.jsonl")}
labels = {r["study"]: r for r in csv.DictReader(open("config/design_labels_v1.tsv"), delimiter="\t")}
preds = {json.loads(l)["study"]: json.loads(l) for l in open(a.pred)}
studies = [s for s in labels if s in preds]
def attr_summary(p):
    out = []
    for k in p["sample_attributes"][:14]:
        lv = k["levels"]; s = ", ".join(f"{kk}({vv})" for kk, vv in list(lv.items())[:8]) if isinstance(lv, dict) else str(lv)
        out.append(f"[{k['key']}] {s[:120]}")
    return "\n".join(out)
def contrasts(p):
    cs = p.get("contrasts") or []
    return "\n".join(f"{i+1}. [{c['comparison_key']}] {c['case_level']} vs {c['control_level']}" + (f" ({c['note']})" if c.get("note") else "") for i, c in enumerate(cs))
rows = []
for s in studies:
    p, l, q = packets[s], labels[s], preds[s]
    agree = (l["label"] == "include") == bool(q["usable_design"])
    rows.append([s, l["organism"], str(p["n_samples"]), str(p.get("tissue_category")), str(p["study_title"]), str(p["abstract"] or "")[:700], attr_summary(p),
                 l["label"], f"[{l['group_key']}] {l['case_level']} vs {l['control_level']}" if l["label"] == "include" else l["reason"],
                 "usable" if q["usable_design"] else "unusable", q["design_type"], q.get("exclusion_reason") or "", contrasts(q), "; ".join(f"{f['action']} {f['key']}={f['value']}" for f in q.get("subset_filters", [])), q.get("sex_composition", ""), ", ".join(q.get("scope_flags", [])), q["confidence"], q["rationale"], "agree" if agree else "disagree", "", "", ""])
if a.sample:
    random.seed(a.seed); dis = [r for r in rows if r[18] == "disagree"]; agr = [r for r in rows if r[18] == "agree"]
    rows = dis + random.sample(agr, max(0, a.sample - len(dis))) if len(agr) > a.sample - len(dis) else rows
wb = Workbook(); ws = wb.active; ws.title = "review"
hdr = ["study", "organism", "n", "tissue", "title", "abstract (700 chars)", "sample attributes (key: levels)", "Claude manual label", "manual label reason/contrast", "LLM verdict", "LLM design type", "LLM exclusion reason", "LLM contrast candidates (ranked)", "LLM subset filters", "LLM sex composition", "LLM scope flags", "LLM confidence", "LLM rationale", "label-LLM agreement", "▶ reviewer verdict (usable / unusable)", "▶ reviewer contrast definition", "▶ reviewer notes"]
ws.append(hdr)
for c in range(1, len(hdr) + 1):
    ws.cell(row=1, column=c).font = Font(bold=True); ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="DDDDDD" if c < 20 else "FFF2CC"); ws.cell(row=1, column=c).alignment = Alignment(wrap_text=True, vertical="top")
for r in rows: ws.append([str(v) if isinstance(v, (list, dict)) else v for v in r])
widths = [11, 8, 5, 10, 40, 60, 70, 10, 40, 9, 22, 16, 45, 25, 12, 16, 8, 50, 8, 16, 30, 30]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    if row[18].value == "disagree":
        for c in row[:19]: c.fill = PatternFill("solid", fgColor="FCE4D6")
ws.freeze_panes = "B2"
wb.save(a.out); print("rows:", len(rows), "→", a.out)
