#!/usr/bin/env python3
"""source_data/*.tsv → per-panel xlsx + combined xlsx (sheet = panel). Usage: python3 27_source_data_xlsx.py [source_data dir]"""
import csv, glob, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
SD = sys.argv[1] if len(sys.argv) > 1 else "results/figures/draft_v3/source_data"
def load(f):
    rows = []
    with open(f, newline="") as fh:
        for row in csv.reader(fh, delimiter="\t"):
            out = []
            for c in row:
                if c in ("", "NA"): out.append(None); continue
                try: out.append(float(c) if any(ch in c for ch in ".eE-") or c.isdigit() else c)
                except ValueError: out.append(c)
            rows.append(out)
    return rows
def fill(ws, rows):
    for r in rows: ws.append(r)
    for c in ws[1]: c.font = Font(bold=True)
    for i, col in enumerate(ws.columns, 1):
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 60)
    ws.freeze_panes = "A2"
combined = Workbook(); combined.remove(combined.active)
files = sorted(glob.glob(os.path.join(SD, "*.tsv")))
for f in files:
    name = os.path.basename(f)[:-4]; rows = load(f)
    wb = Workbook(); fill(wb.active, rows); wb.active.title = name[:31]; wb.save(os.path.join(SD, name + ".xlsx"))
    fill(combined.create_sheet(name[:31]), rows)
combined.save(os.path.join(SD, "SourceData_all_panels.xlsx"))
print(len(files), "panel xlsx +", "SourceData_all_panels.xlsx")
